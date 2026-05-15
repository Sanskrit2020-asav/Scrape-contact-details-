"""Extract plain text from prepared itinerary files (PDF, DOCX, TXT/MD)."""

from __future__ import annotations

from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            # A single unreadable page should not sink the whole document.
            continue
    return "\n".join(parts)


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_text(path: Path) -> str:
    """Return the document's text, or '' if the format is unsupported/unreadable."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _read_pdf(path)
        if suffix == ".docx":
            return _read_docx(path)
        if suffix in (".txt", ".md"):
            return _read_text(path)
    except Exception:
        return ""
    return ""
