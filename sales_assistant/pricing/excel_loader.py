"""Read the company's Excel (.xlsx) cost estimator into a :class:`RateCard`.

Design goals (per the brief):
  * Read all pricing data and preserve every rule we can map.
  * Extract formulas where possible (openpyxl exposes them) and record them so
    the spreadsheet's logic is auditable, not silently re-implemented.
  * Never invent prices: anything the sheet doesn't contain keeps the documented
    default, and the loader reports exactly which keys it populated.

The expected (but flexible) layout is a two-column "key, value" rates sheet
and/or a permits sheet. Because every studio's spreadsheet differs, the loader
is tolerant: it scans all sheets for ``label: number`` pairs and matches labels
to known rate fields by fuzzy keyword. Unmapped pairs are returned as ``extras``
so nothing is lost — wire them in by extending ``_FIELD_KEYWORDS``.

openpyxl is an optional dependency; if it is absent the loader raises a clear
error and the platform falls back to default rates.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ..models import SourceRef, SourceType
from .rates import RateCard

# Map RateCard scalar fields → keywords that might label them in the sheet.
_FIELD_KEYWORDS: dict[str, tuple[str, ...]] = {
    "guide_per_day": ("guide per day", "guide/day", "guide day", "guide rate"),
    "porter_per_day": ("porter per day", "porter/day", "porter day", "porter rate"),
    "climbing_guide_per_day": ("climbing guide", "sherpa guide", "climbing sherpa"),
    "teahouse_pppn": ("teahouse", "tea house", "lodge", "trek accommodation"),
    "hotel_3star_pppn": ("3 star", "3-star", "three star", "city hotel", "standard hotel"),
    "hotel_luxury_pppn": ("luxury hotel", "5 star", "5-star", "deluxe hotel"),
    "meals_per_day": ("meal", "food per day", "full board", "board"),
    "domestic_flight_per_person": ("flight", "lukla", "domestic air", "airfare"),
    "private_vehicle_per_day": ("private vehicle", "jeep", "car per day", "private car"),
    "tourist_bus_per_person": ("tourist bus", "bus per person", "coach"),
    "airport_transfer_flat": ("airport transfer", "pick up", "pickup", "drop"),
    "misc_per_person": ("misc", "miscellaneous", "sundry", "buffer"),
    "default_markup_pct": ("markup", "margin", "profit %", "profit percent"),
}

_NUM_RE = re.compile(r"-?\d[\d,]*\.?\d*")


@dataclass
class ExcelLoadResult:
    rate_card: RateCard
    populated_fields: list[str] = field(default_factory=list)
    permits_loaded: int = 0
    activities_loaded: int = 0
    formulas: dict[str, str] = field(default_factory=dict)   # "Sheet!A1" -> "=B1*C1"
    extras: dict[str, float] = field(default_factory=dict)   # unmapped label -> value
    warnings: list[str] = field(default_factory=list)

    def summary(self) -> str:
        return (
            f"Loaded {len(self.populated_fields)} rate fields, "
            f"{self.permits_loaded} permits, {self.activities_loaded} activities, "
            f"{len(self.formulas)} formulas captured, {len(self.extras)} unmapped values."
        )


def _to_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    m = _NUM_RE.search(str(value))
    if not m:
        return None
    try:
        return float(m.group(0).replace(",", ""))
    except ValueError:
        return None


def _match_field(label: str) -> str | None:
    low = label.strip().lower()
    if not low:
        return None
    for field_name, keywords in _FIELD_KEYWORDS.items():
        if any(kw in low for kw in keywords):
            return field_name
    return None


def _looks_like_permit(label: str) -> bool:
    low = label.lower()
    return any(w in low for w in ("permit", "tims", "acap", "card", "conservation", "national park", "rural municipality"))


def _looks_like_activity(label: str) -> bool:
    low = label.lower()
    return any(w in low for w in ("safari", "flight tour", "scenic", "rafting", "paragliding", "city tour", "helicopter", "tour"))


def load_rate_card_from_excel(path: str | Path, base: RateCard | None = None) -> ExcelLoadResult:
    """Parse an .xlsx estimator into a RateCard, preserving formulas and extras."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise RuntimeError(
            "openpyxl is required to read .xlsx pricing files. Install with "
            "`pip install openpyxl`."
        ) from exc

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    card = base or RateCard()
    result = ExcelLoadResult(rate_card=card)
    result.rate_card.source = SourceRef(SourceType.EXCEL, path.name, str(path))

    # Two passes: cached values (data_only) and formulas, iterated in lock-step.
    wb_values = load_workbook(path, data_only=True, read_only=True)
    wb_formulas = load_workbook(path, data_only=False, read_only=True)

    from itertools import zip_longest

    for ws in wb_values.worksheets:
        fws = wb_formulas[ws.title] if ws.title in wb_formulas.sheetnames else None
        frows = fws.iter_rows() if fws is not None else iter(())
        for row, frow in zip_longest(ws.iter_rows(), frows, fillvalue=()):
            frow = list(frow) if frow else []
            # Find the first text cell (label) and first numeric cell (value).
            label = None
            value = None
            value_idx = None
            for idx, cell in enumerate(row):
                if cell.value is None:
                    continue
                if label is None and isinstance(cell.value, str) and not _NUM_RE.fullmatch(cell.value.strip()):
                    label = cell.value.strip()
                    continue
                num = _to_number(cell.value)
                if num is not None and value is None:
                    value, value_idx = num, idx

            # Capture formulas for auditability — at the value column, or anywhere
            # in the row when the cached value is missing (file never opened in Excel).
            if label and frow:
                if value_idx is not None and value_idx < len(frow):
                    fval = frow[value_idx].value
                    if isinstance(fval, str) and fval.startswith("="):
                        result.formulas[f"{ws.title}!{frow[value_idx].coordinate}"] = fval
                elif value is None:
                    for fcell in frow:
                        if isinstance(fcell.value, str) and fcell.value.startswith("="):
                            result.formulas[f"{ws.title}!{fcell.coordinate}"] = fcell.value
                            break

            if not label or value is None:
                continue

            # Permit/activity classification wins over the broad rate keywords
            # (e.g. "Everest scenic flight" is an activity, not the flight rate).
            if _looks_like_permit(label):
                card.permit_costs[label] = value
                result.permits_loaded += 1
                continue
            if _looks_like_activity(label):
                card.activity_costs[label.lower()] = value
                result.activities_loaded += 1
                continue
            field_name = _match_field(label)
            if field_name:
                setattr(card, field_name, value)
                result.populated_fields.append(field_name)
            else:
                result.extras[label] = value

    wb_values.close()
    wb_formulas.close()

    if not result.populated_fields and not result.permits_loaded:
        result.warnings.append(
            "No recognisable rate labels were found. Check the sheet layout or extend "
            "_FIELD_KEYWORDS in pricing/excel_loader.py."
        )
    return result
